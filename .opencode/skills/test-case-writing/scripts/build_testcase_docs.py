#!/usr/bin/python3
"""
Build a Word (.docx) test-case document and an Excel (.xlsx) test-case workbook
from ONE JSON file (schema: references/json-schema.md).

Usage:
    python build_testcase_docs.py testcases.json --out-dir OUT [--basename NAME]
           [--docx-only | --xlsx-only]

Both outputs are generated from the same data so they never drift apart.
Requires: python-docx, openpyxl
"""
import argparse
import json
import os
import sys
from collections import OrderedDict

# --------------------------------------------------------------------------- #
# Labels (vi / en)
# --------------------------------------------------------------------------- #
L = {
    "vi": {
        "doc_info": "Thông tin tài liệu", "doc_id": "Mã tài liệu", "project": "Dự án / Hệ thống",
        "feature": "Chức năng", "version": "Phiên bản", "date": "Ngày", "author": "Người viết",
        "reviewer": "Người duyệt", "sut_version": "Phiên bản phần mềm được test",
        "environment": "Môi trường", "intro": "Giới thiệu", "scope": "Phạm vi",
        "references": "Tài liệu tham chiếu", "assumptions": "Giả định",
        "conventions": "Quy ước đặt tên và trình bày",
        "conv_text": [
            "Mã test case: TC-<MÔ-ĐUN>-<số thứ tự 2–3 chữ số>, ví dụ TC-LOGIN-003.",
            "Tiêu đề theo cú pháp Hành động + Chức năng + Điều kiện (Action + Function + Condition).",
            "Mỗi test case một mục tiêu, dữ liệu cụ thể, kết quả mong đợi kiểm chứng được.",
            "Trạng thái: Not run / Pass / Fail / Blocked / Skipped.",
        ],
        "requirements": "Yêu cầu và ma trận truy vết", "req_id": "Mã yêu cầu", "req_text": "Nội dung yêu cầu",
        "covered_by": "Test case bao phủ", "not_covered": "CHƯA bao phủ",
        "summary": "Tổng hợp test case", "suite": "Bộ test", "format": "Định dạng",
        "fmt_steps": "Từng bước (step-by-step)", "fmt_matrix": "Ma trận dữ liệu (data matrix)",
        "details": "Chi tiết test case", "id_title": "Mã / Tiêu đề", "purpose": "Mục đích",
        "trace": "Truy vết yêu cầu", "priority": "Ưu tiên", "type": "Loại", "method": "Phương pháp",
        "setup": "Thiết lập / Tiền điều kiện", "test_data": "Dữ liệu test", "steps": "Các bước",
        "step_no": "#", "action": "Hành động", "expected": "Kết quả mong đợi",
        "expected_result": "Kết quả mong đợi (tổng)", "cleanup": "Dọn dẹp (self-cleaning)",
        "status_bug": "Trạng thái / Mã lỗi", "actual": "Kết quả thực tế", "status": "Trạng thái",
        "bug_id": "Mã lỗi", "notes": "Ghi chú", "origin": "Nguồn gốc",
        "procedure": "Quy trình chung (áp dụng cho mọi dòng)", "covers": "Bao phủ",
        "matrix_note": "Mỗi dòng = một test case. Quy trình giống nhau, chỉ dữ liệu thay đổi.",
        "review": "Kết quả review test case", "review_source": "Đối tượng review",
        "scores": "Điểm theo 7 thuộc tính (0–2)", "total": "Tổng /14", "verdict": "Kết luận",
        "defects": "Lỗi phát hiện", "category": "Loại lỗi", "severity": "Mức độ",
        "description": "Mô tả", "fix": "Cách sửa", "checklist": "Checklist review",
        "check_item": "Tiêu chí", "result": "Kết quả", "review_summary": "Nhận xét chung",
        "ai_verdict": "Đối chiếu kết quả review của AI", "confirmed": "Lỗi xác nhận",
        "false_pos": "AI báo sai (false positive)", "missed": "AI bỏ sót",
        "refactor": "Đối chiếu refactor", "old": "Bản cũ", "new": "Test case mới", "change": "Thay đổi",
        "ai_decl": "Khai báo sử dụng AI", "ai_tool": "Công cụ AI", "ai_prompts": "Prompt chính",
        "ai_changes": "Những gì đã sửa so với bản AI và lý do", "ai_verify": "Cách kiểm chứng",
        "cover": "Tổng quan", "tc_sheet": "Test Cases", "trace_sheet": "Truy vết",
        "review_sheet": "Review", "checklist_sheet": "Checklist",
        "count_total": "Tổng số test case", "count_status": "Theo trạng thái",
        "yes": "Có", "no": "Không", "partial": "Một phần",
        "attrs": OrderedDict([
            ("accurate", "Chính xác"), ("economical", "Tiết kiệm"), ("repeatable", "Lặp lại được"),
            ("traceable", "Truy vết được"), ("appropriate", "Phù hợp"),
            ("self_standing", "Độc lập"), ("self_cleaning", "Tự dọn dẹp"),
        ]),
        "default_checklist": [
            "Đạt 7 thuộc tính: chính xác, tiết kiệm, lặp lại được, truy vết được, phù hợp, độc lập, tự dọn dẹp",
            "Có mã/tên và mục đích gắn với yêu cầu",
            "Mô tả đủ thiết lập: môi trường, dữ liệu, tiền điều kiện, quyền truy cập",
            "Mọi hành động đều có kết quả mong đợi; ghi bằng chứng khi cần",
            "Câu chủ động; ≤ 15 bước; ma trận chạy ≤ 20 phút",
            "Thứ tự đúng theo kịch bản nghiệp vụ so với các test khác",
            "Theo đúng quy ước đặt tên/đánh số; có phiên bản ứng với phần mềm được test",
            "Được lưu trữ có kiểm soát truy cập và sao lưu (quản lý cấu hình)",
        ],
    },
    "en": {
        "doc_info": "Document information", "doc_id": "Document ID", "project": "Project / System",
        "feature": "Feature", "version": "Version", "date": "Date", "author": "Author",
        "reviewer": "Reviewer", "sut_version": "Software version under test",
        "environment": "Environment", "intro": "Introduction", "scope": "Scope",
        "references": "References", "assumptions": "Assumptions",
        "conventions": "Naming and presentation conventions",
        "conv_text": [
            "Test case ID: TC-<MODULE>-<2–3 digit sequence>, e.g. TC-LOGIN-003.",
            "Titles follow Action + Function + Condition.",
            "One case = one objective, concrete data, and a checkable expected result.",
            "Status: Not run / Pass / Fail / Blocked / Skipped.",
        ],
        "requirements": "Requirements and traceability matrix", "req_id": "Req. ID",
        "req_text": "Requirement", "covered_by": "Covered by", "not_covered": "NOT covered",
        "summary": "Test case summary", "suite": "Suite", "format": "Format",
        "fmt_steps": "Step-by-step", "fmt_matrix": "Data matrix",
        "details": "Test case details", "id_title": "ID / Title", "purpose": "Purpose",
        "trace": "Traces to", "priority": "Priority", "type": "Type", "method": "Method",
        "setup": "Setup / Preconditions", "test_data": "Test data", "steps": "Steps",
        "step_no": "#", "action": "Action", "expected": "Expected result",
        "expected_result": "Expected result (overall)", "cleanup": "Cleanup (self-cleaning)",
        "status_bug": "Status / Bug ID", "actual": "Actual result", "status": "Status",
        "bug_id": "Bug ID", "notes": "Notes", "origin": "Origin",
        "procedure": "Common procedure (applies to every row)", "covers": "Covers",
        "matrix_note": "One row = one test case. Same procedure every time; only the data varies.",
        "review": "Test case review results", "review_source": "Reviewed artifact",
        "scores": "Scores on the 7 attributes (0–2)", "total": "Total /14", "verdict": "Verdict",
        "defects": "Defects found", "category": "Category", "severity": "Severity",
        "description": "Description", "fix": "Fix", "checklist": "Review checklist",
        "check_item": "Criterion", "result": "Result", "review_summary": "Overall comments",
        "ai_verdict": "Verifying the AI reviewer", "confirmed": "Confirmed defects",
        "false_pos": "AI false positives", "missed": "Missed by AI",
        "refactor": "Refactoring map", "old": "Legacy item", "new": "New test case(s)", "change": "Change",
        "ai_decl": "AI-use declaration", "ai_tool": "AI tool", "ai_prompts": "Key prompts",
        "ai_changes": "What I changed from the AI output and why", "ai_verify": "How I verified",
        "cover": "Overview", "tc_sheet": "Test Cases", "trace_sheet": "Traceability",
        "review_sheet": "Review", "checklist_sheet": "Checklist",
        "count_total": "Total test cases", "count_status": "By status",
        "yes": "Yes", "no": "No", "partial": "Partial",
        "attrs": OrderedDict([
            ("accurate", "Accurate"), ("economical", "Economical"), ("repeatable", "Repeatable"),
            ("traceable", "Traceable"), ("appropriate", "Appropriate"),
            ("self_standing", "Self-standing"), ("self_cleaning", "Self-cleaning"),
        ]),
        "default_checklist": [
            "Accurate, economical, repeatable, traceable, appropriate, self-standing, self-cleaning",
            "Has a name/number and a stated purpose tied to a requirement",
            "Specifies setup: environment, data, prerequisites, access",
            "Every action has an expected result; proofs noted where needed",
            "Active language; ≤ 15 steps; matrix runs ≤ 20 minutes",
            "Correct business-scenario order relative to other tests",
            "Follows naming/numbering conventions; versioned to the software under test",
            "Stored with controlled access and backup (configuration management)",
        ],
    },
}
STATUSES = ["Not run", "Pass", "Fail", "Blocked", "Skipped"]

PRIMARY = "1F4E79"      # dark blue header
LIGHT = "DCE6F1"        # light blue label cells
ZEBRA = "F5F8FC"
FAIL_RED = "F8CBAD"
PASS_GREEN = "C6EFCE"
WARN_YELLOW = "FFEB9C"


def as_list(v):
    if v is None or v == "":
        return []
    return v if isinstance(v, list) else [v]


def joined(v, sep="\n"):
    return sep.join(str(x) for x in as_list(v))


def all_cases(data):
    for s in data.get("suites", []):
        for c in s.get("cases", []):
            yield s, c


# --------------------------------------------------------------------------- #
# WORD
# --------------------------------------------------------------------------- #
def build_docx(data, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    meta = data.get("meta", {})
    lang = meta.get("language", "vi")
    T = L["en" if lang == "en" else "vi"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor.from_string(PRIMARY)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run((meta.get("doc_id", "") + "  ·  ") if meta.get("doc_id") else "")
    run.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = fp.add_run()
        r.font.size = Pt(8)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)

    def no_split(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        el.set(qn("w:val"), "true")
        trPr.append(el)

    def set_text(cell, text, bold=False, color=None, size=None, italic=False):
        cell.text = ""
        lines = str(text if text is not None else "").split("\n")
        p = cell.paragraphs[0]
        for i, line in enumerate(lines):
            if i:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            r = p.add_run(line)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def widths(table, cms):
        table.autofit = False
        tblPr = table._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for i, w in enumerate(cms):
            if i < len(table.columns):
                table.columns[i].width = Cm(w)
        for row in table.rows:
            for i, w in enumerate(cms):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)

    def grid(headers, rows, cms, zebra=True, font=9.5):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            set_text(c, h, bold=True, color="FFFFFF", size=font)
            shade(c, PRIMARY)
        repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                set_text(cells[i], v, size=font)
                if zebra and ri % 2 == 1:
                    shade(cells[i], ZEBRA)
        widths(t, cms)
        doc.add_paragraph()
        return t

    def kv_table(pairs, cms=(4.5, 12.5)):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in pairs:
            cells = t.add_row().cells
            set_text(cells[0], k, bold=True, size=9.5)
            shade(cells[0], LIGHT)
            set_text(cells[1], v, size=9.5)
        widths(t, cms)
        doc.add_paragraph()
        return t

    def bullets(items):
        for it in as_list(items):
            doc.add_paragraph(str(it), style="List Bullet")

    # ---------------- Title ----------------
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run(meta.get("title", "Test Case Document"))
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(PRIMARY)
    if meta.get("feature") or meta.get("project"):
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = sp.add_run(" — ".join(x for x in [meta.get("project"), meta.get("feature")] if x))
        rr.font.size = Pt(12)
        rr.italic = True

    doc.add_heading(T["doc_info"], level=1)
    info = [(T[k], meta.get(k)) for k in
            ("doc_id", "project", "feature", "version", "date", "author", "reviewer", "sut_version", "environment")
            if meta.get(k)]
    kv_table(info)

    # ---------------- 1. Intro ----------------
    sec_no = 1
    doc.add_heading(f"{sec_no}. {T['intro']}", level=1)
    if meta.get("scope"):
        doc.add_heading(f"{sec_no}.1 {T['scope']}", level=2)
        doc.add_paragraph(meta["scope"])
    if meta.get("references"):
        doc.add_heading(f"{sec_no}.2 {T['references']}", level=2)
        bullets(meta["references"])
    if meta.get("assumptions"):
        doc.add_heading(f"{sec_no}.3 {T['assumptions']}", level=2)
        bullets(meta["assumptions"])
    doc.add_heading(f"{sec_no}.4 {T['conventions']}", level=2)
    bullets(T["conv_text"])

    # ---------------- 2. Requirements & traceability ----------------
    reqs = data.get("requirements", [])
    if reqs:
        sec_no += 1
        doc.add_heading(f"{sec_no}. {T['requirements']}", level=1)
        cover = {r["id"]: [] for r in reqs}
        for _, c in all_cases(data):
            for rid in as_list(c.get("requirements")):
                cover.setdefault(rid, []).append(c["id"])
        rows = []
        for r_ in reqs:
            ids = cover.get(r_["id"], [])
            rows.append([r_["id"], r_.get("text", ""), ", ".join(ids) if ids else "⚠ " + T["not_covered"]])
        grid([T["req_id"], T["req_text"], T["covered_by"]], rows, (2.6, 9.4, 5.0))

    # ---------------- 3. Summary ----------------
    sec_no += 1
    doc.add_heading(f"{sec_no}. {T['summary']}", level=1)
    rows = []
    for s, c in all_cases(data):
        rows.append([c.get("id", ""), c.get("title", ""), ", ".join(as_list(c.get("requirements"))),
                     c.get("priority", ""), c.get("type", ""), s.get("id", "")])
    grid(["ID", T["id_title"].split("/")[-1].strip(), T["trace"], T["priority"], T["type"], T["suite"]],
         rows, (2.6, 7.6, 2.4, 1.6, 1.6, 1.2), font=9)

    # ---------------- 4. Details ----------------
    sec_no += 1
    doc.add_heading(f"{sec_no}. {T['details']}", level=1)
    for si, s in enumerate(data.get("suites", []), 1):
        fmt = s.get("format", "steps")
        doc.add_heading(f"{sec_no}.{si} {s.get('id', '')} — {s.get('name', '')}", level=2)
        if s.get("description"):
            doc.add_paragraph(s["description"])
        meta_line = f"{T['format']}: {T['fmt_matrix'] if fmt == 'matrix' else T['fmt_steps']}"
        p = doc.add_paragraph()
        p.add_run(meta_line).italic = True

        if fmt == "matrix":
            pairs = []
            if s.get("preconditions"):
                pairs.append((T["setup"], joined(s["preconditions"])))
            if pairs:
                kv_table(pairs)
            if s.get("procedure"):
                doc.add_heading(T["procedure"], level=3)
                grid([T["step_no"], T["action"], T["expected"]],
                     [[i, st_.get("action", ""), st_.get("expected", "")] for i, st_ in
                      enumerate(s["procedure"], 1)], (1.0, 8.0, 8.0))
            cols = s.get("columns", [])
            headers = ["ID"] + cols + [T["expected"], T["covers"], T["trace"], T["status"]]
            n_in = len(cols)
            in_w = max(1.3, min(2.4, 6.5 / max(1, n_in)))
            cms = [2.4] + [in_w] * n_in + [4.2, 2.6, 1.9, 1.6]
            rows = []
            for c in s.get("cases", []):
                inp = c.get("inputs", {})
                rows.append([c.get("id", "")] + [inp.get(col, "") for col in cols] +
                            [c.get("expected_result", ""), c.get("covers", c.get("title", "")),
                             ", ".join(as_list(c.get("requirements"))), c.get("status", "Not run")])
            grid(headers, rows, cms, font=9)
            p = doc.add_paragraph()
            p.add_run(T["matrix_note"]).italic = True
            if s.get("cleanup"):
                kv_table([(T["cleanup"], joined(s["cleanup"]))])
            # titles list for matrix rows (titles still matter)
            if any(c.get("title") for c in s.get("cases", [])):
                for c in s.get("cases", []):
                    if c.get("title"):
                        bp = doc.add_paragraph(style="List Bullet")
                        bp.add_run(c["id"] + " — ").bold = True
                        bp.add_run(c["title"])
            continue

        for c in s.get("cases", []):
            doc.add_heading(f"{c.get('id', '')} — {c.get('title', '')}", level=3)
            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"

            def row(k, v):
                cells = t.add_row().cells
                set_text(cells[0], k, bold=True, size=9.5)
                shade(cells[0], LIGHT)
                set_text(cells[1], v, size=9.5)
                no_split(t.rows[-1])

            row(T["id_title"], f"{c.get('id', '')} — {c.get('title', '')}")
            if c.get("purpose"):
                row(T["purpose"], c["purpose"])
            if c.get("requirements"):
                row(T["trace"], ", ".join(as_list(c["requirements"])))
            attrs = " · ".join(x for x in [
                f"{T['priority']}: {c['priority']}" if c.get("priority") else "",
                f"{T['type']}: {c['type']}" if c.get("type") else "",
                f"{T['method']}: {c['method']}" if c.get("method") else "",
                f"{T['origin']}: {c['origin']}" if c.get("origin") else ""] if x)
            if attrs:
                row(T["priority"] + " / " + T["type"], attrs)
            if c.get("preconditions"):
                row(T["setup"], joined(c["preconditions"]))
            if c.get("test_data"):
                row(T["test_data"], joined(c["test_data"]))
            widths(t, (4.2, 12.8))

            steps = c.get("steps", [])
            if steps:
                st_t = doc.add_table(rows=1, cols=3)
                st_t.style = "Table Grid"
                for i, h in enumerate([T["step_no"], T["action"], T["expected"]]):
                    set_text(st_t.rows[0].cells[i], h, bold=True, color="FFFFFF", size=9.5)
                    shade(st_t.rows[0].cells[i], PRIMARY)
                repeat_header(st_t.rows[0])
                for i, step in enumerate(steps, 1):
                    cells = st_t.add_row().cells
                    set_text(cells[0], str(i), size=9.5)
                    set_text(cells[1], step.get("action", ""), size=9.5)
                    set_text(cells[2], step.get("expected", ""), size=9.5)
                    no_split(st_t.rows[-1])
                widths(st_t, (1.0, 8.0, 8.0))

            t2 = doc.add_table(rows=0, cols=2)
            t2.style = "Table Grid"

            def row2(k, v):
                cells = t2.add_row().cells
                set_text(cells[0], k, bold=True, size=9.5)
                shade(cells[0], LIGHT)
                set_text(cells[1], v, size=9.5)
                no_split(t2.rows[-1])

            if c.get("expected_result"):
                row2(T["expected_result"], c["expected_result"])
            if c.get("cleanup"):
                row2(T["cleanup"], joined(c["cleanup"]))
            stv = c.get("status", "Not run")
            boxes = "   ".join(("☒ " if stv == s_ else "☐ ") + s_ for s_ in ["Pass", "Fail", "Blocked", "Skipped"])
            row2(T["status_bug"], f"{boxes}      {T['bug_id']}: {c.get('bug_id') or '______'}")
            if c.get("actual_result"):
                row2(T["actual"], c["actual_result"])
            if c.get("notes"):
                row2(T["notes"], c["notes"])
            widths(t2, (4.2, 12.8))
            doc.add_paragraph()

    # ---------------- Review ----------------
    rv = data.get("review")
    if rv:
        sec_no += 1
        doc.add_heading(f"{sec_no}. {T['review']}", level=1)
        if rv.get("source"):
            kv_table([(T["review_source"], rv["source"])])
        if rv.get("summary"):
            doc.add_paragraph(rv["summary"])
        items = rv.get("items", [])
        if items:
            doc.add_heading(T["scores"], level=2)
            keys = list(T["attrs"].keys())
            heads = ["ID"] + [T["attrs"][k] for k in keys] + [T["total"], T["verdict"]]
            rows = []
            for it in items:
                sc = it.get("scores", {})
                vals = [sc.get(k, "") for k in keys]
                tot = sum(v for v in vals if isinstance(v, (int, float)))
                rows.append([it.get("case_id", "")] + vals + [tot, it.get("verdict", "")])
            grid(heads, rows, [2.4] + [1.45] * 7 + [1.3, 2.2], font=8.5)
            drows = []
            for it in items:
                for d in it.get("defects", []):
                    drows.append([it.get("case_id", ""), d.get("category", ""), d.get("severity", ""),
                                  d.get("description", ""), d.get("fix", "")])
            if drows:
                doc.add_heading(T["defects"], level=2)
                grid(["ID", T["category"], T["severity"], T["description"], T["fix"]], drows,
                     (2.2, 2.6, 1.6, 5.3, 5.3), font=9)
        av = rv.get("ai_review_check")
        if av:
            doc.add_heading(T["ai_verdict"], level=2)
            cols = [as_list(av.get("confirmed")), as_list(av.get("false_positives")), as_list(av.get("missed"))]
            n = max(len(x) for x in cols) if any(cols) else 0
            rows = [[(cols[j][i] if i < len(cols[j]) else "") for j in range(3)] for i in range(n)]
            grid([T["confirmed"], T["false_pos"], T["missed"]], rows, (5.6, 5.7, 5.7), font=9)
        cl = rv.get("checklist")
        if cl:
            doc.add_heading(T["checklist"], level=2)
            grid([T["check_item"], T["result"], T["notes"]],
                 [[x.get("item", ""), x.get("result", ""), x.get("note", "")] for x in cl], (9.0, 2.2, 5.8), font=9)

    # ---------------- Refactor map ----------------
    rf = data.get("refactor")
    if rf and rf.get("mapping"):
        sec_no += 1
        doc.add_heading(f"{sec_no}. {T['refactor']}", level=1)
        if rf.get("source"):
            kv_table([(T["review_source"], rf["source"])])
        grid([T["old"], T["new"], T["change"]],
             [[m.get("old", ""), ", ".join(as_list(m.get("new"))), m.get("change", "")] for m in rf["mapping"]],
             (5.5, 3.0, 8.5), font=9)

    # ---------------- AI declaration ----------------
    ad = data.get("ai_declaration")
    if ad:
        sec_no += 1
        doc.add_heading(f"{sec_no}. {T['ai_decl']}", level=1)
        kv_table([(T["ai_tool"], ad.get("tool", "")),
                  (T["ai_prompts"], joined(ad.get("prompts"), "\n\n")),
                  (T["ai_changes"], joined(ad.get("changes"))),
                  (T["ai_verify"], ad.get("verification", ""))])

    doc.save(path)


# --------------------------------------------------------------------------- #
# EXCEL
# --------------------------------------------------------------------------- #
def build_xlsx(data, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    meta = data.get("meta", {})
    lang = meta.get("language", "vi")
    T = L["en" if lang == "en" else "vi"]

    wb = Workbook()
    thin = Side(style="thin", color="A6A6A6")
    thick = Side(style="medium", color=PRIMARY)
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    HFILL = PatternFill("solid", fgColor=PRIMARY)
    LFILL = PatternFill("solid", fgColor=LIGHT)
    ZFILL = PatternFill("solid", fgColor=ZEBRA)
    HFONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    BFONT = Font(name="Calibri", size=10)
    WRAP = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")

    def header(ws, row, heads, widths_=None):
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.fill, c.font, c.border = HFILL, HFONT, BORDER
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths_:
            for i, w in enumerate(widths_, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[row].height = 30

    def cell(ws, r, c, v, fill=None, bold=False, center=False):
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name="Calibri", size=10, bold=bold)
        x.border = BORDER
        x.alignment = CENTER if center else WRAP
        if fill:
            x.fill = fill
        return x

    def status_formatting(ws, rng):
        for val, color in (("Pass", PASS_GREEN), ("Fail", FAIL_RED), ("Blocked", WARN_YELLOW)):
            ws.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))

    def status_validation(ws, rng):
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)

    def est_height(texts, widths_):
        lines = 1
        for t, w in zip(texts, widths_):
            if t is None:
                continue
            s = str(t)
            n = sum(max(1, -(-len(part) // max(1, int(w * 1.1)))) for part in s.split("\n"))
            lines = max(lines, n)
        return min(400, 14 * lines + 4)

    # ---------- Sheet: Test Cases (step-by-step suites) ----------
    ws = wb.active
    ws.title = T["tc_sheet"]
    heads = ["ID", "Title" if lang == "en" else "Tiêu đề", T["trace"], T["priority"], T["type"],
             T["setup"], T["test_data"], T["step_no"], T["action"], T["expected"],
             T["cleanup"], T["actual"], T["status"], T["bug_id"], T["notes"], T["suite"]]
    W = [13, 34, 11, 9, 11, 28, 18, 4, 36, 36, 22, 20, 10, 10, 18, 8]
    header(ws, 1, heads, W)
    ws.freeze_panes = "C2"
    r = 2
    status_cells = []
    case_idx = 0
    for s in data.get("suites", []):
        if s.get("format", "steps") == "matrix":
            continue
        for c in s.get("cases", []):
            steps = c.get("steps") or [{"action": "", "expected": c.get("expected_result", "")}]
            start = r
            fill = ZFILL if case_idx % 2 else None
            case_idx += 1
            for i, step in enumerate(steps, 1):
                exp = step.get("expected", "")
                if i == len(steps) and c.get("expected_result") and c.get("steps"):
                    exp = (exp + "\n" if exp else "") + "⇒ " + c["expected_result"]
                vals = [c.get("id", ""), c.get("title", ""), ", ".join(as_list(c.get("requirements"))),
                        c.get("priority", ""), c.get("type", ""), joined(c.get("preconditions")),
                        joined(c.get("test_data")), i, step.get("action", ""), exp,
                        joined(c.get("cleanup")), c.get("actual_result", ""), c.get("status", "Not run"),
                        c.get("bug_id", ""), c.get("notes", ""), s.get("id", "")]
                for col, v in enumerate(vals, 1):
                    cell(ws, r, col, v, fill=fill, bold=(col == 1), center=col in (4, 8, 13, 16))
                ws.row_dimensions[r].height = est_height([step.get("action"), exp], [W[8], W[9]])
                r += 1
            end = r - 1
            # merge case-level columns across its step rows (ID repeated in col A is kept merged too)
            if end > start:
                for col in (1, 2, 3, 4, 5, 6, 7, 11, 12, 13, 14, 15, 16):
                    ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            # thick separator under each case
            for col in range(1, len(heads) + 1):
                b = ws.cell(row=end, column=col).border
                ws.cell(row=end, column=col).border = Border(left=b.left, right=b.right, top=b.top, bottom=thick)
            status_cells.append(f"M{start}")
    last_tc_row = max(2, r - 1)
    status_validation(ws, f"M2:M{last_tc_row}")
    status_formatting(ws, f"M2:M{last_tc_row}")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{last_tc_row}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- Sheets: matrix suites ----------
    matrix_status_ranges = []
    for s in data.get("suites", []):
        if s.get("format") != "matrix":
            continue
        name = (s.get("id") or "Matrix")[:31]
        mws = wb.create_sheet(name)
        mws.cell(row=1, column=1, value=f"{s.get('id', '')} — {s.get('name', '')}").font = Font(
            bold=True, size=13, color=PRIMARY)
        row = 2
        if s.get("description"):
            mws.cell(row=row, column=1, value=s["description"]).alignment = WRAP
            row += 1
        if s.get("preconditions"):
            mws.cell(row=row, column=1, value=T["setup"]).font = Font(bold=True)
            mws.cell(row=row, column=2, value=joined(s["preconditions"], "; ")).alignment = WRAP
            row += 1
        if s.get("procedure"):
            mws.cell(row=row, column=1, value=T["procedure"]).font = Font(bold=True, color=PRIMARY)
            row += 1
            for i, stp in enumerate(s["procedure"], 1):
                mws.cell(row=row, column=1, value=f"{i}. {stp.get('action', '')}")
                if stp.get("expected"):
                    mws.cell(row=row, column=1).value += f"  →  {stp['expected']}"
                row += 1
        if s.get("cleanup"):
            mws.cell(row=row, column=1, value=T["cleanup"] + ": " + joined(s["cleanup"], "; ")).font = Font(italic=True)
            row += 1
        row += 1
        cols = s.get("columns", [])
        heads_m = ["ID", "Title" if lang == "en" else "Tiêu đề"] + cols + [
            T["expected"], T["covers"], T["trace"], T["priority"], T["actual"], T["status"], T["bug_id"]]
        widths_m = [13, 34] + [11] * len(cols) + [30, 18, 11, 9, 20, 10, 10]
        header(mws, row, heads_m, widths_m)
        hdr_row = row
        mws.freeze_panes = mws.cell(row=row + 1, column=3)
        row += 1
        st_col = len(heads_m) - 1
        for k, c in enumerate(s.get("cases", [])):
            fill = ZFILL if k % 2 else None
            inp = c.get("inputs", {})
            vals = [c.get("id", ""), c.get("title", "")] + [inp.get(col, "") for col in cols] + [
                c.get("expected_result", ""), c.get("covers", ""), ", ".join(as_list(c.get("requirements"))),
                c.get("priority", ""), c.get("actual_result", ""), c.get("status", "Not run"), c.get("bug_id", "")]
            for j, v in enumerate(vals, 1):
                cell(mws, row, j, v, fill=fill, bold=(j == 1), center=(3 <= j < 3 + len(cols)) or j == st_col)
            mws.row_dimensions[row].height = est_height([c.get("title"), c.get("expected_result")], [34, 30])
            row += 1
        L_ = get_column_letter(st_col)
        a_, b_ = hdr_row + 1, max(hdr_row + 1, row - 1)
        rng = f"{L_}{a_}:{L_}{b_}"
        status_validation(mws, rng)
        status_formatting(mws, rng)
        matrix_status_ranges.append(f"'{name}'!${L_}${a_}:${L_}${b_}")
        mws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(heads_m))}{max(hdr_row, row - 1)}"
        mws.page_setup.orientation = "landscape"
        mws.sheet_properties.pageSetUpPr.fitToPage = True
        mws.page_setup.fitToHeight = 0

    # ---------- Sheet: Traceability ----------
    reqs = data.get("requirements", [])
    if reqs:
        tws = wb.create_sheet(T["trace_sheet"])
        ids = [c.get("id") for _, c in all_cases(data)]
        heads_t = [T["req_id"], T["req_text"], "#TC"] + ids
        header(tws, 1, heads_t, [12, 45, 6] + [5] * len(ids))
        for j in range(4, len(heads_t) + 1):
            tws.cell(row=1, column=j).alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom")
        tws.row_dimensions[1].height = 90
        case_reqs = {c.get("id"): set(as_list(c.get("requirements"))) for _, c in all_cases(data)}
        for i, rq in enumerate(reqs, 2):
            cell(tws, i, 1, rq["id"], bold=True)
            cell(tws, i, 2, rq.get("text", ""))
            last = get_column_letter(len(heads_t))
            cell(tws, i, 3, f'=COUNTIF(D{i}:{last}{i},"X")', center=True)
            for j, cid in enumerate(ids, 4):
                cell(tws, i, j, "X" if rq["id"] in case_reqs.get(cid, set()) else "", center=True)
            tws.row_dimensions[i].height = est_height([rq.get("text")], [45])
        n = len(reqs) + 1
        tws.conditional_formatting.add(f"C2:C{n}", CellIsRule(operator="equal", formula=["0"],
                                                               fill=PatternFill("solid", fgColor=FAIL_RED)))
        if ids:
            tws.conditional_formatting.add(
                f"D2:{get_column_letter(len(heads_t))}{n}",
                CellIsRule(operator="equal", formula=['"X"'], fill=PatternFill("solid", fgColor=LIGHT)))
        tws.freeze_panes = "D2"

    # ---------- Sheet: Review ----------
    rv = data.get("review")
    if rv and rv.get("items"):
        rws = wb.create_sheet(T["review_sheet"])
        keys = list(T["attrs"].keys())
        heads_r = ["ID"] + [T["attrs"][k] for k in keys] + [T["total"], T["verdict"],
                                                            T["category"], T["severity"], T["description"], T["fix"]]
        header(rws, 1, heads_r, [13] + [10] * 7 + [8, 10, 16, 9, 40, 40])
        r = 2
        for it in rv["items"]:
            defects = it.get("defects") or [{}]
            start = r
            for d in defects:
                sc = it.get("scores", {})
                for j, k in enumerate(keys, 2):
                    cell(rws, r, j, sc.get(k, ""), center=True)
                cell(rws, r, 1, it.get("case_id", ""), bold=True)
                cell(rws, r, 9, f"=SUM(B{start}:H{start})", center=True)
                cell(rws, r, 10, it.get("verdict", ""), center=True)
                cell(rws, r, 11, d.get("category", ""))
                cell(rws, r, 12, d.get("severity", ""), center=True)
                cell(rws, r, 13, d.get("description", ""))
                cell(rws, r, 14, d.get("fix", ""))
                rws.row_dimensions[r].height = est_height([d.get("description"), d.get("fix")], [40, 40])
                r += 1
            if r - 1 > start:
                for col in range(1, 11):
                    rws.merge_cells(start_row=start, start_column=col, end_row=r - 1, end_column=col)
        rng = f"B2:H{r - 1}"
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True)
        rws.add_data_validation(dv)
        dv.add(rng)
        for val, color in (("0", FAIL_RED), ("1", WARN_YELLOW), ("2", PASS_GREEN)):
            rws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[val],
                                                           fill=PatternFill("solid", fgColor=color)))
        rws.freeze_panes = "B2"
        av = rv.get("ai_review_check")
        if av:
            r += 2
            rws.cell(row=r, column=1, value=T["ai_verdict"]).font = Font(bold=True, size=12, color=PRIMARY)
            r += 1
            heads_av = [T["confirmed"], T["false_pos"], T["missed"]]
            colsav = [13, 14, 11]  # wide columns
            for j, h in enumerate(heads_av):
                x = rws.cell(row=r, column=colsav[j], value=h)
                x.fill, x.font, x.border = HFILL, HFONT, BORDER
            lists = [as_list(av.get("confirmed")), as_list(av.get("false_positives")), as_list(av.get("missed"))]
            for i in range(max(len(x) for x in lists) if any(lists) else 0):
                r += 1
                for j in range(3):
                    cell(rws, r, colsav[j], lists[j][i] if i < len(lists[j]) else "")

    # ---------- Sheet: Checklist ----------
    cws = wb.create_sheet(T["checklist_sheet"])
    header(cws, 1, ["#", T["check_item"], T["result"], T["notes"]], [4, 70, 12, 40])
    given = (rv or {}).get("checklist") or [{"item": x} for x in T["default_checklist"]]
    for i, it in enumerate(given, 2):
        cell(cws, i, 1, i - 1, center=True)
        cell(cws, i, 2, it.get("item", ""))
        cell(cws, i, 3, it.get("result", ""), center=True)
        cell(cws, i, 4, it.get("note", ""))
        cws.row_dimensions[i].height = est_height([it.get("item"), it.get("note")], [70, 40])
    dv = DataValidation(type="list", formula1=f'"{T["yes"]},{T["no"]},{T["partial"]}"', allow_blank=True)
    cws.add_data_validation(dv)
    dv.add(f"C2:C{len(given) + 1}")

    # ---------- Sheet: Overview (first) ----------
    ows = wb.create_sheet(T["cover"], 0)
    ows.column_dimensions["A"].width = 30
    ows.column_dimensions["B"].width = 60
    ows.cell(row=1, column=1, value=meta.get("title", "Test Cases")).font = Font(bold=True, size=16, color=PRIMARY)
    r = 3
    for k in ("doc_id", "project", "feature", "version", "date", "author", "reviewer", "sut_version", "environment"):
        if meta.get(k):
            cell(ows, r, 1, T[k], fill=LFILL, bold=True)
            cell(ows, r, 2, meta[k])
            r += 1
    r += 1
    ows.cell(row=r, column=1, value=T["count_status"]).font = Font(bold=True, size=12, color=PRIMARY)
    r += 1
    tc_name = T["tc_sheet"]
    srcs = [f"'{tc_name}'!$M$2:$M${last_tc_row}"] + matrix_status_ranges
    first = r
    for stv in STATUSES:
        cell(ows, r, 1, stv, fill=LFILL, bold=True)
        cell(ows, r, 2, "=" + "+".join(f'COUNTIF({x},"{stv}")' for x in srcs), center=False)
        r += 1
    cell(ows, r, 1, T["count_total"], fill=LFILL, bold=True)
    cell(ows, r, 2, f"=SUM(B{first}:B{r - 1})", bold=True)
    r += 2
    ows.cell(row=r, column=1, value=T["summary"]).font = Font(bold=True, size=12, color=PRIMARY)
    r += 1
    for s in data.get("suites", []):
        cell(ows, r, 1, s.get("id", ""), fill=LFILL, bold=True)
        cell(ows, r, 2, f"{s.get('name', '')} — {T['fmt_matrix'] if s.get('format') == 'matrix' else T['fmt_steps']}"
                        f" — {len(s.get('cases', []))} TC")
        r += 1
    r += 1
    ows.cell(row=r, column=1, value=T["conventions"]).font = Font(bold=True, size=12, color=PRIMARY)
    for line in T["conv_text"]:
        r += 1
        ows.cell(row=r, column=1, value="• " + line)

    for w_ in wb.worksheets:
        w_.sheet_properties.pageSetUpPr.fitToPage = True
        w_.page_setup.fitToWidth = 1
        w_.page_setup.fitToHeight = 0
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json_path")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--basename", default=None)
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--docx-only", action="store_true")
    g.add_argument("--xlsx-only", action="store_true")
    a = ap.parse_args()
    with open(a.json_path, encoding="utf-8") as f:
        data = json.load(f)
    base = a.basename or os.path.splitext(os.path.basename(a.json_path))[0]
    os.makedirs(a.out_dir, exist_ok=True)
    outs = []
    if not a.xlsx_only:
        p = os.path.join(a.out_dir, base + ".docx")
        build_docx(data, p)
        outs.append(p)
    if not a.docx_only:
        p = os.path.join(a.out_dir, base + ".xlsx")
        build_xlsx(data, p)
        outs.append(p)
    for p in outs:
        print("wrote", p)


if __name__ == "__main__":
    sys.exit(main())
